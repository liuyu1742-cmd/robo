import hashlib
import json
import unittest
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\sjtu101\Desktop\3.docx")
OUTPUT = ROOT / "outputs" / "3_项目测试更新版.docx"
BENCHMARK = ROOT / "outputs" / "decision_planning_benchmark_20260731.json"
SOURCE_SHA256 = "F1DCB6E240B74A3A23FB9DB47C33C5EF85CFDDD3DAC0EFA4CA223C2F3C75E52C"


class DecisionTestDocxOutputTest(unittest.TestCase):
    def test_source_was_not_overwritten(self):
        digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest().upper()
        self.assertEqual(digest, SOURCE_SHA256)

    def test_updated_document_contains_exact_case_and_result_headers(self):
        document = Document(OUTPUT)

        self.assertEqual(len(document.tables), 3)
        self.assertEqual(
            [cell.text for cell in document.tables[0].rows[0].cells],
            ["用例 ID", "测试场景工况", "测试操作步骤", "预期输出结果", "合格判定标准"],
        )
        self.assertEqual(
            [cell.text for cell in document.tables[1].rows[0].cells],
            ["用例 ID", "测试指令", "10次耗时（s）", "平均耗时（s）", "真实输出结果", "合格判定"],
        )
        self.assertEqual(
            [row.cells[0].text for row in document.tables[0].rows[1:]],
            ["DEC-001", "DEC-002", "DEC-003"],
        )
        self.assertEqual(
            [row.cells[0].text for row in document.tables[1].rows[1:]],
            ["DEC-001", "DEC-002", "DEC-003"],
        )
        self.assertEqual(
            [cell.text for cell in document.tables[2].rows[0].cells],
            [
                "证据 ID",
                "用例/轮次",
                "开始时间",
                "结束时间",
                "计时器原始值与耗时",
                "输出 SHA-256",
            ],
        )
        self.assertEqual(len(document.tables[2].rows), 32)

    def test_result_values_match_raw_benchmark(self):
        report = json.loads(BENCHMARK.read_text(encoding="utf-8"))
        document = Document(OUTPUT)
        result_table = document.tables[1]

        for row, case in zip(result_table.rows[1:], report["cases"]):
            self.assertEqual(row.cells[1].text, case["instruction"])
            expected_times = "\n".join(
                f"{index}. {value:.9f}"
                for index, value in enumerate(case["elapsed_seconds"], start=1)
            )
            self.assertEqual(row.cells[2].text, expected_times)
            self.assertEqual(row.cells[3].text, f"{case['average_seconds']:.9f}")
            self.assertIn(case["real_output"], row.cells[4].text)
            self.assertIn("合格", row.cells[5].text)
        self.assertTrue(report["all_qualified"])

    def test_each_benchmark_run_has_a_documented_evidence_row(self):
        report = json.loads(BENCHMARK.read_text(encoding="utf-8"))
        document = Document(OUTPUT)
        evidence_rows = {
            row.cells[0].text: [cell.text for cell in row.cells]
            for row in document.tables[2].rows[1:]
        }

        for case in report["cases"]:
            for run in case["runs"]:
                cells = evidence_rows[run["evidence_id"]]
                self.assertEqual(cells[1], f"{case['case_id']}/{run['run_index']:02d}")
                self.assertIn(str(run["started_perf_counter_ns"]), cells[4])
                self.assertIn(str(run["ended_perf_counter_ns"]), cells[4])
                self.assertIn(run["output_sha256"], cells[5])
        replan = report["cases"][-1]["replanning"]
        self.assertIn(replan["evidence_id"], evidence_rows)

    def test_report_embeds_six_visual_evidence_images_with_explanations(self):
        document = Document(OUTPUT)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)

        self.assertEqual(len(document.inline_shapes), 6)
        self.assertIn("3.2.6.6 图像化测试证据与结果解释", text)
        for title in (
            "图1 测试执行总览",
            "图2 VLA用例10次运行证据",
            "图3 多设备协同10次运行及重规划证据",
            "图4 三个既有动作解析子任务调用结果",
            "图5 10次耗时对比与合格阈值",
            "图6 计时代码与真实输出对照",
        ):
            self.assertIn(title, text)
        self.assertIn("图片数据来源", text)
        self.assertIn("elapsed_seconds=(ended_perf_counter_ns-started_perf_counter_ns)", text)
        self.assertIn("软件决策规划测试，不代表实体家电已经执行", text)


if __name__ == "__main__":
    unittest.main()
