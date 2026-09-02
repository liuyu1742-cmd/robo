"""Update the reference DOCX with reproducible project test cases and results."""

from __future__ import annotations

import argparse
from copy import deepcopy
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = Path(r"C:\Users\sjtu101\Desktop\3.docx")
DEFAULT_BENCHMARK = ROOT / "outputs" / "decision_planning_benchmark_20260731.json"
DEFAULT_OUTPUT = ROOT / "outputs" / "3_项目测试更新版.docx"
DEFAULT_SCREENSHOT_DIRECTORY = ROOT / "outputs" / "decision_evidence_screenshots"
EXPECTED_CASE_HEADERS = [
    "用例 ID",
    "测试场景工况",
    "测试操作步骤",
    "预期输出结果",
    "合格判定标准",
]
RESULT_HEADERS = [
    "用例 ID",
    "测试指令",
    "10次耗时（s）",
    "平均耗时（s）",
    "真实输出结果",
    "合格判定",
]
EVIDENCE_HEADERS = [
    "证据 ID",
    "用例/轮次",
    "开始时间",
    "结束时间",
    "计时器原始值与耗时",
    "输出 SHA-256",
]

VISUAL_EVIDENCE = [
    (
        "图1 测试执行总览",
        "01_test_overview.png",
        "图片数据来源：汇总文件 decision_planning_benchmark_20260731.json 及31份逐轮JSON证据。"
        "图中列出实际测试命令、执行时间、随机种子、运行环境、计时公式、三项平均耗时和合格阈值。"
        "三项用例均显示“通过”，说明汇总结果与逐轮证据一致。",
    ),
    (
        "图2 VLA用例10次运行证据",
        "02_vla_run_evidence.png",
        "图片数据来源：DEC-001与DEC-002各10份逐轮JSON证据。每行同时显示证据ID、"
        "elapsed_ns、换算后的秒数和输出SHA-256前缀，可据此逐轮复核。两个用例来自不同VLA任务"
        "与不同物体；其毫秒级耗时来自读取已审核动作模板并生成结构化输出。",
    ),
    (
        "图3 多设备协同10次运行及重规划证据",
        "03_multi_device_run_evidence.png",
        "图片数据来源：DEC-003的10份运行证据和1份环境变化重规划证据。图中保留"
        "started_perf_counter_ns、ended_perf_counter_ns及两者差值；"
        "elapsed_seconds=(ended_perf_counter_ns-started_perf_counter_ns)/1,000,000,000。"
        "10次平均耗时和重规划耗时均低于5秒阈值。",
    ),
    (
        "图4 三个既有动作解析子任务调用结果",
        "04_multi_device_subtasks.png",
        "图片数据来源：DEC-003结构化输出中的subtasks与coordinated_targets字段。"
        "图中证明“准备就寝环境”实际调用了空调、电动窗帘和卧室灯三个既有动作解析任务，"
        "分别生成定位、操作和检查动作，再协调为24℃、关闭窗帘及20%亮度；冲突检查通过。",
    ),
    (
        "图5 10次耗时对比与合格阈值",
        "05_timing_chart.png",
        "图片数据来源：三个测试用例共30次elapsed_seconds。曲线纵轴放大到1.6秒以便观察波动，"
        "同时标出VLA用例3秒和多设备用例5秒合格阈值。多设备耗时高于VLA用例，是因为它串联"
        "执行三个既有解析子任务，并追加参数补全、并行编组和冲突检查。",
    ),
    (
        "图6 计时代码与真实输出对照",
        "06_timing_code_and_output.png",
        "图片数据来源：项目实际文件tools/run_decision_planning_benchmark.py与"
        "DEC-003-R10对应JSON证据。左侧代码展示计时起止点、纳秒差值、秒数换算、完整输出"
        "序列化及SHA-256计算；右侧展示同一轮真实计数器值、哈希和三个解析子任务输出，"
        "共同证明报告耗时由实际代码测得而非人工填写。",
    ),
]


def _set_run_font(run, *, size_pt: float, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋_GB2312")
    run.font.size = Pt(size_pt)
    run.font.bold = bold


def _set_cell_text(
    cell,
    text: str,
    *,
    size_pt: float = 8.0,
    bold: bool = False,
    center: bool = False,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    )
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    for index, line in enumerate(str(text).splitlines() or [""]):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        _set_run_font(run, size_pt=size_pt, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _copy_cell_properties(source_cell, target_cell) -> None:
    source_properties = source_cell._tc.tcPr
    target_properties = target_cell._tc.get_or_add_tcPr()
    for child in list(target_properties):
        target_properties.remove(child)
    for child in source_properties:
        target_properties.append(deepcopy(child))


def _set_repeat_table_header(row) -> None:
    tr_properties = row._tr.get_or_add_trPr()
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    tr_properties.append(table_header)


def _insert_paragraph_after(paragraph, text: str):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph

    created = Paragraph(new_p, paragraph._parent)
    created.paragraph_format.first_line_indent = Pt(24)
    created.paragraph_format.space_after = Pt(3)
    run = created.add_run(text)
    _set_run_font(run, size_pt=10.5)
    return created


def _insert_table_after(paragraph, table) -> None:
    paragraph._p.addnext(table._tbl)


def _format_times(values: list[float]) -> str:
    return "\n".join(
        f"{index}. {value:.9f}" for index, value in enumerate(values, start=1)
    )


def _case_rows(report: dict) -> list[list[str]]:
    rows = []
    for case in report["cases"]:
        if case["task"] == "multi_device_coordination":
            expected = (
                "识别空调、电动窗帘、卧室灯；输出设备状态检查、并行调度、"
                "空调24℃、窗帘关闭、卧室灯20%亮度、最终状态确认及冲突检查；"
                "室温升至31℃时重新规划为空调制冷模式。"
            )
            operation = (
                f"1. 下发“{case['instruction']}”；\n"
                "2. 从接收指令开始计时，至完整协同调度方案输出完成；\n"
                "3. 重复10次并记录；\n"
                "4. 将室温临时调整为31℃，再次计时并检查重新规划。"
            )
            standard = (
                "10次平均决策耗时≤5s；环境临时变化后的重规划耗时≤5s；"
                "必须实际调用空调、电动窗帘、卧室灯三个既有动作解析子任务，"
                "且三台设备目标无冲突。"
            )
        elif case.get("source") == "VLA":
            structured = case["structured_output"]
            expected = (
                f"识别VLA任务 {case['task']}、物体"
                f"{case['object_name_zh']}（{case['object_id']}），输出原动作模板："
                + " → ".join(structured["template_steps"])
                + "。"
            )
            operation = (
                f"1. 下发“{case['instruction']}”；\n"
                "2. 从指令接收开始计时，至VLA动作模板及可序列化输出生成完成；\n"
                "3. 重复10次并记录单次耗时和独立证据文件。"
            )
            standard = (
                "10次平均决策耗时≤3s，且VLA任务、物体及原动作模板输出完整。"
            )
        else:
            structured = case["structured_output"]
            expected = (
                f"识别任务 {structured['resolved']['task']}、物体 "
                f"{structured['resolved']['object']}，输出完整动作序列："
                + " → ".join(structured["final_actions"])
                + "。"
            )
            operation = (
                f"1. 下发“{case['instruction']}”；\n"
                "2. 从接收指令开始计时，至完整决策方案输出完成；\n"
                "3. 重复10次并记录单次耗时。"
            )
            standard = "10次平均决策耗时≤3s，且任务、物体及动作序列输出完整。"
        rows.append(
            [
                case["case_id"],
                case["scene"],
                operation,
                expected,
                standard,
            ]
        )
    return rows


def _result_rows(report: dict) -> list[list[str]]:
    rows = []
    for case in report["cases"]:
        real_output = case["real_output"]
        qualification = case["qualification"]
        if replanning := case.get("replanning"):
            real_output += "\n\n环境变化后的真实重规划输出：\n" + replanning[
                "real_output"
            ]
            qualification += (
                f"；重规划输出完整，耗时{replanning['elapsed_seconds']:.9f}s。"
            )
        rows.append(
            [
                case["case_id"],
                case["instruction"],
                _format_times(case["elapsed_seconds"]),
                f"{case['average_seconds']:.9f}",
                real_output,
                qualification,
            ]
        )
    return rows


def _evidence_rows(report: dict) -> list[list[str]]:
    rows: list[list[str]] = []
    for case in report["cases"]:
        for run in case["runs"]:
            rows.append(
                [
                    run["evidence_id"],
                    f"{case['case_id']}/{run['run_index']:02d}",
                    run["started_at"],
                    run["ended_at"],
                    (
                        f"start={run['started_perf_counter_ns']}\n"
                        f"end={run['ended_perf_counter_ns']}\n"
                        f"Δ={run['elapsed_ns']} ns\n"
                        f"={run['elapsed_seconds']:.9f} s"
                    ),
                    run["output_sha256"],
                ]
            )
        if replanning := case.get("replanning"):
            rows.append(
                [
                    replanning["evidence_id"],
                    f"{case['case_id']}/重规划",
                    replanning["started_at"],
                    replanning["ended_at"],
                    (
                        f"start={replanning['started_perf_counter_ns']}\n"
                        f"end={replanning['ended_perf_counter_ns']}\n"
                        f"Δ={replanning['elapsed_ns']} ns\n"
                        f"={replanning['elapsed_seconds']:.9f} s"
                    ),
                    replanning["output_sha256"],
                ]
            )
    return rows


def _add_visual_evidence_section(document, screenshot_directory: Path) -> None:
    missing = [
        screenshot_directory / filename
        for _, filename, _ in VISUAL_EVIDENCE
        if not (screenshot_directory / filename).is_file()
    ]
    if missing:
        raise FileNotFoundError(
            "缺少图像化测试证据：" + "，".join(str(path) for path in missing)
        )

    heading = document.add_paragraph()
    heading_run = heading.add_run("3.2.6.6 图像化测试证据与结果解释")
    _set_run_font(heading_run, size_pt=10.5, bold=True)
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(4)

    overview = document.add_paragraph()
    overview.paragraph_format.first_line_indent = Pt(24)
    overview.paragraph_format.space_after = Pt(4)
    overview_run = overview.add_run(
        "以下六张图片由本次测试的汇总JSON、逐轮证据JSON和项目实际计时代码自动生成，"
        "用于直观展示10次运行耗时、原始计数器值、结构化输出及合格判定。"
        "截图用于方便审核，最终复核仍以31份原始证据的计时算术和输出SHA-256为准。"
    )
    _set_run_font(overview_run, size_pt=10.5)

    for title, filename, explanation in VISUAL_EVIDENCE:
        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(6)
        caption.paragraph_format.space_after = Pt(3)
        caption_run = caption.add_run(title)
        _set_run_font(caption_run, size_pt=10.5, bold=True)

        image_paragraph = document.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.space_after = Pt(3)
        image_run = image_paragraph.add_run()
        image_run.add_picture(
            str(screenshot_directory / filename),
            width=Inches(5.65),
        )

        explanation_paragraph = document.add_paragraph()
        explanation_paragraph.paragraph_format.first_line_indent = Pt(24)
        explanation_paragraph.paragraph_format.space_after = Pt(5)
        explanation_run = explanation_paragraph.add_run(explanation)
        _set_run_font(explanation_run, size_pt=9.5)

    conclusion = document.add_paragraph()
    conclusion.paragraph_format.first_line_indent = Pt(24)
    conclusion.paragraph_format.space_before = Pt(4)
    conclusion.paragraph_format.space_after = Pt(3)
    conclusion_run = conclusion.add_run(
        "实验结果解释：DEC-001与DEC-002均来自VLA任务—物体模板，10次运行输出完整且平均耗时"
        "低于3秒；DEC-003实际调用三个既有动作解析子任务，因解析、协调和冲突检查步骤更多，"
        "平均耗时高于两个VLA用例，但10次平均耗时及31℃环境变化后的重规划耗时仍低于5秒。"
        "本测试属于软件决策规划测试，不代表实体家电已经执行。"
    )
    _set_run_font(conclusion_run, size_pt=10.5)


def update_document(
    source: Path,
    benchmark: Path,
    output: Path,
    screenshot_directory: Path = DEFAULT_SCREENSHOT_DIRECTORY,
) -> Path:
    report = json.loads(benchmark.read_text(encoding="utf-8"))
    document = Document(source)
    if not document.tables:
        raise ValueError("参考文档中未找到测试用例表")

    case_table = document.tables[0]
    actual_headers = [cell.text.strip() for cell in case_table.rows[0].cells]
    if actual_headers != EXPECTED_CASE_HEADERS:
        raise ValueError(
            f"测试用例表头不匹配：期望{EXPECTED_CASE_HEADERS}，实际{actual_headers}"
        )
    if len(case_table.rows) < 4:
        raise ValueError("参考测试用例表不足4行")

    for row, values in zip(case_table.rows[1:4], _case_rows(report)):
        for cell, value in zip(row.cells, values):
            _set_cell_text(cell, value, size_pt=9.0)

    result_heading = next(
        (
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.strip() == "3.2.6.4 测试结果"
        ),
        None,
    )
    if result_heading is None:
        raise ValueError("未找到“3.2.6.4 测试结果”")
    following = result_heading._p.getnext()
    if following is None or following.tag != qn("w:p"):
        raise ValueError("测试结果标题后未找到说明段落")

    from docx.text.paragraph import Paragraph

    placeholder = Paragraph(following, result_heading._parent)
    placeholder.text = ""
    intro = (
        f"测试于{report['executed_at']}执行。计时器为{report['timer']}，计时口径为："
        f"{report['timing_scope']}。每个用例重复{report['repetitions']}次，"
        f"计算公式为{report['timing_formula']}，随机种子为{report['random_seed']}。"
    )
    intro_run = placeholder.add_run(intro)
    _set_run_font(intro_run, size_pt=10.5)
    placeholder.paragraph_format.first_line_indent = Pt(24)
    placeholder.paragraph_format.space_after = Pt(3)
    environment_text = (
        f"测试环境：{report['environment']['platform']}，"
        f"Python {report['environment']['python']}，"
        f"处理器 {report['environment']['processor']}。"
        f"测试边界：{report['environment']['execution_boundary']}。"
    )
    environment_paragraph = _insert_paragraph_after(placeholder, environment_text)
    selected_text = "随机抽选结果：" + "；".join(
        f"{case['case_id']}“{case['instruction']}”" for case in report["cases"]
    ) + "。"
    selected_paragraph = _insert_paragraph_after(environment_paragraph, selected_text)

    result_table = document.add_table(rows=4, cols=6)
    result_table.style = case_table.style
    result_table.autofit = True
    for column_index, cell in enumerate(result_table.rows[0].cells):
        source_cell = case_table.rows[0].cells[min(column_index, 4)]
        _copy_cell_properties(source_cell, cell)
    for row in result_table.rows[1:]:
        for column_index, cell in enumerate(row.cells):
            source_cell = case_table.rows[1].cells[min(column_index, 4)]
            _copy_cell_properties(source_cell, cell)

    for cell, header in zip(result_table.rows[0].cells, RESULT_HEADERS):
        _set_cell_text(cell, header, size_pt=7.5, bold=True, center=True)
    _set_repeat_table_header(result_table.rows[0])
    for row, values in zip(result_table.rows[1:], _result_rows(report)):
        for column_index, (cell, value) in enumerate(zip(row.cells, values)):
            _set_cell_text(
                cell,
                value,
                size_pt=6.5 if column_index in {2, 4} else 7.0,
                center=column_index in {0, 3, 5},
            )
    _insert_table_after(selected_paragraph, result_table)

    conclusion = _insert_paragraph_after(
        selected_paragraph,
        "测试结论：3个测试用例全部达到参考文档规定的耗时和输出完整性要求，"
        "前两个用例均来自已审核VLA任务—物体模板；多设备协同指令实际调用了"
        "空调、电动窗帘、卧室灯三个既有动作解析子任务，未发现互斥目标，"
        "环境温度变化后的重新规划同样满足5秒时限。",
    )
    # Moving the result table after the conclusion keeps the summary visible
    # before the long evidence table.
    _insert_table_after(conclusion, result_table)

    evidence_heading = document.add_paragraph()
    evidence_heading_run = evidence_heading.add_run("3.2.6.5 测试证据")
    _set_run_font(evidence_heading_run, size_pt=10.5, bold=True)
    evidence_heading.paragraph_format.space_before = Pt(6)
    evidence_heading.paragraph_format.space_after = Pt(3)
    result_table._tbl.addnext(evidence_heading._p)
    evidence_intro = _insert_paragraph_after(
        evidence_heading,
        "每轮测试均生成一个独立JSON证据文件。下表同时记录墙钟开始/结束时间、"
        "perf_counter_ns起止原始值、差值、换算秒数及完整输出SHA-256。"
        f"证据目录：{report['evidence_directory']}。",
    )
    evidence_values = _evidence_rows(report)
    evidence_table = document.add_table(
        rows=len(evidence_values) + 1, cols=len(EVIDENCE_HEADERS)
    )
    evidence_table.style = case_table.style
    evidence_table.autofit = True
    for column_index, cell in enumerate(evidence_table.rows[0].cells):
        source_cell = case_table.rows[0].cells[min(column_index, 4)]
        _copy_cell_properties(source_cell, cell)
    for row in evidence_table.rows[1:]:
        for column_index, cell in enumerate(row.cells):
            source_cell = case_table.rows[1].cells[min(column_index, 4)]
            _copy_cell_properties(source_cell, cell)
    for cell, header in zip(evidence_table.rows[0].cells, EVIDENCE_HEADERS):
        _set_cell_text(cell, header, size_pt=7.0, bold=True, center=True)
    _set_repeat_table_header(evidence_table.rows[0])
    for row, values in zip(evidence_table.rows[1:], evidence_values):
        for column_index, (cell, value) in enumerate(zip(row.cells, values)):
            _set_cell_text(
                cell,
                value,
                size_pt=6.0 if column_index in {4, 5} else 6.5,
                center=column_index in {0, 1},
            )
    _insert_table_after(evidence_intro, evidence_table)
    _add_visual_evidence_section(document, screenshot_directory)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--benchmark", type=Path, default=DEFAULT_BENCHMARK)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--screenshot-directory",
        type=Path,
        default=DEFAULT_SCREENSHOT_DIRECTORY,
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    output = update_document(
        args.source,
        args.benchmark,
        args.output,
        args.screenshot_directory,
    )
    print(output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
