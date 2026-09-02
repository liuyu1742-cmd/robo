from docx import Document
from pathlib import Path

path = Path("3_测试报告补充.docx")
doc = Document(path)
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text:
        print(f"P{i}: style={p.style.name!r} text={text}")
for i, table in enumerate(doc.tables):
    print(f"TABLE {i}: {len(table.rows)}x{len(table.columns)}")
    for row in table.rows[:6]:
        print(" | ".join(cell.text.replace("\n", " / ") for cell in row.cells))
