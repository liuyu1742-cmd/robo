"""Refine the test-case table and add screenshot context to Desktop/3.docx."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


DOC_PATH = Path(r"C:\Users\sjtu101\Desktop\3.docx")
REPORT_PATH = Path(r"C:\RobotProject\RobotProject\datasets\manual_instruction_ten_case_actual_report.json")

CASE_DETAILS = {
    "ACT-001": {
        "scenario": "书房桌面收纳",
        "steps": "1. 输入文本指令：把书桌上的文件和文具整理整齐。\n2. 系统定位书桌及桌面待整理物品，解析收纳意图。\n3. 读取匹配任务、物体、操作和原始动作序列。\n4. 与书桌收纳标准序列逐项比对。",
        "expected": "匹配：organizing / desk / organize。\n标准动作序列：\n1. locate(desk)\n2. grasp(desk)\n3. move(storage)\n4. place(desk)",
        "criterion": "任务、物体和操作均匹配为书桌收纳；定位、抓取、移动、放置四步完整，顺序正确且无对象替换。",
    },
    "ACT-002": {
        "scenario": "卧室阅读灯关闭",
        "steps": "1. 输入文本指令：睡前请关闭阅读灯并确认已经熄灭。\n2. 系统解析关闭设备意图，并定位阅读灯。\n3. 记录原始模型生成的开关控制动作。\n4. 比对是否输出关灯动作及状态检查步骤。",
        "expected": "匹配：appliance_management / reading_lamp / turn_off。\n标准动作序列：\n1. locate(reading_lamp)\n2. turn_off(reading_lamp)\n3. inspect(reading_lamp)",
        "criterion": "必须输出 turn_off(reading_lamp)，不得输出 turn_on；完成定位、关灯和熄灭状态检查后判定合格。",
    },
    "ACT-003": {
        "scenario": "出门门锁安防检查",
        "steps": "1. 输入文本指令：出门前检查智能门锁是否已经锁好。\n2. 系统识别安防检查工况并定位智能门锁。\n3. 读取门锁控制及状态核验动作。\n4. 对比锁定和检查步骤是否齐全。",
        "expected": "匹配：security_monitoring / smart_lock / lock。\n标准动作序列：\n1. locate(smart_lock)\n2. lock(smart_lock)\n3. inspect(smart_lock)",
        "criterion": "智能门锁匹配正确；锁定动作在检查动作之前，且定位、锁定、核验三项均存在。",
    },
    "ACT-004": {
        "scenario": "客厅遥控器取送",
        "steps": "1. 输入文本指令：请把茶几上的遥控器递给我。\n2. 系统识别取送对象为遥控器、目标为用户。\n3. 记录抓取、递送和释放动作。\n4. 对比对象、交接目标和动作顺序。",
        "expected": "匹配：object_fetching / remote_control / fetch。\n标准动作序列：\n1. locate(remote_control)\n2. grasp(remote_control)\n3. move(person)\n4. release(remote_control)",
        "criterion": "遥控器与用户交接目标匹配正确；抓取后移动至 person，再释放遥控器，四步不得缺失或倒序。",
    },
    "ACT-005": {
        "scenario": "卫生间马桶清洁",
        "steps": "1. 输入文本指令：请清洁马桶内外表面。\n2. 系统识别卫浴清洁对象为马桶。\n3. 记录刷洗、冲洗和清洁结果检查动作。\n4. 对比清洁流程的完整性。",
        "expected": "匹配：cleaning / toilet / clean。\n标准动作序列：\n1. locate(toilet)\n2. scrub(toilet)\n3. rinse(toilet)\n4. inspect(toilet)",
        "criterion": "马桶对象匹配正确；刷洗、冲洗、检查三类关键动作齐全，且均作用于 toilet。",
    },
    "ACT-006": {
        "scenario": "废旧电池分类回收",
        "steps": "1. 输入文本指令：把门口的废电池放入专用回收盒。\n2. 系统识别电池为需分类处理物品。\n3. 记录定位、抓取、移至垃圾箱和处置动作。\n4. 检查处置对象和目标容器是否正确。",
        "expected": "匹配：waste_disposal / battery / dispose。\n标准动作序列：\n1. locate(battery)\n2. grasp(battery)\n3. move(trash_bin)\n4. dispose(battery)",
        "criterion": "电池识别和 dispose 操作正确；必须先抓取电池、再移动至 trash_bin 后处置，无错误物体或遗漏步骤。",
    },
    "ACT-007": {
        "scenario": "早餐电饭煲烹饪",
        "steps": "1. 输入文本指令：烹饪前请启动电饭煲煮饭。\n2. 系统识别烹饪设备为电饭煲。\n3. 记录设备启动、烹饪和结束关闭动作。\n4. 对比烹饪流程顺序。",
        "expected": "匹配：smart_cooking / rice_cooker / cook。\n标准动作序列：\n1. locate(rice_cooker)\n2. turn_on(rice_cooker)\n3. cook(rice_cooker)\n4. turn_off(rice_cooker)",
        "criterion": "电饭煲和 cook 操作匹配正确；开机、烹饪、关机依次出现，且动作对象均为 rice_cooker。",
    },
    "ACT-008": {
        "scenario": "玄关可视门铃巡检",
        "steps": "1. 输入文本指令：检查玄关的可视门铃是否正常。\n2. 系统定位可视门铃并识别巡检意图。\n3. 记录设备检查相关原始动作。\n4. 比对定位和巡检动作是否完整。",
        "expected": "匹配：security_monitoring / video_doorbell / inspect。\n标准动作序列：\n1. locate(video_doorbell)\n2. inspect(video_doorbell)",
        "criterion": "可视门铃对象匹配正确；输出应包含定位和 inspect 两步，且不应替换为其他安防设备。",
    },
    "ACT-009": {
        "scenario": "客厅玩具收纳",
        "steps": "1. 输入文本指令：把茶几上的玩具放进收纳箱。\n2. 系统应将玩具识别为待操作对象、收纳箱识别为目标容器。\n3. 记录对象匹配和原始动作序列。\n4. 对比是否围绕 toy 生成收纳动作。",
        "expected": "匹配：organizing / toy / organize。\n标准动作序列：\n1. locate(toy)\n2. grasp(toy)\n3. move(storage)\n4. place(toy)",
        "criterion": "待操作对象必须为 toy，收纳箱只能作为容器；若将 storage_box 识别为操作对象，则判定不合格。",
    },
    "ACT-010": {
        "scenario": "玄关拐杖取送",
        "steps": "1. 输入文本指令：把门口的拐杖拿给我。\n2. 系统识别取送对象为拐杖、交接目标为用户。\n3. 记录定位、抓取、递送与释放动作。\n4. 对比交接对象和动作顺序。",
        "expected": "匹配：object_fetching / walking_cane / fetch。\n标准动作序列：\n1. locate(walking_cane)\n2. grasp(walking_cane)\n3. move(person)\n4. release(walking_cane)",
        "criterion": "拐杖与 person 目标匹配正确；完整生成定位、抓取、递送、释放四步后判定合格。",
    },
}


def shade(cell, fill: str) -> None:
    pr = cell._tc.get_or_add_tcPr()
    item = OxmlElement("w:shd")
    item.set(qn("w:fill"), fill)
    pr.append(item)


def width(cell, value: float) -> None:
    pr = cell._tc.get_or_add_tcPr()
    item = pr.find(qn("w:tcW")) or OxmlElement("w:tcW")
    item.set(qn("w:w"), str(int(value * 1440)))
    item.set(qn("w:type"), "dxa")
    if item.getparent() is None:
        pr.append(item)


def write_cell(cell, text: str, *, size: float = 8.3, bold: bool = False, center: bool = False) -> None:
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            run.font.size = Pt(size)
            run.bold = bold


def add_text(doc, text: str, *, bold: bool = False, size: float = 10.5, placeholder: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if placeholder:
        run.font.color.rgb = __import__("docx.shared", fromlist=["RGBColor"]).RGBColor(192, 0, 0)


def remove_paragraph(paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def action_lines(actions: list[str]) -> str:
    return "\n".join(f"{i}. {a}" for i, a in enumerate(actions, 1))


def main() -> None:
    report = json.loads(REPORT_PATH.read_text(encoding="utf-8"))
    doc = Document(DOC_PATH)
    table = doc.tables[0]
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)

    headers = ["用例 ID", "测试场景工况", "测试操作步骤", "预期输出结果", "合格判定标准"]
    for index, text in enumerate(headers):
        write_cell(table.rows[0].cells[index], text, size=9, bold=True, center=True)
        shade(table.rows[0].cells[index], "D9EAF7")

    for case in report["cases"]:
        detail = CASE_DETAILS[case["case_id"]]
        cells = table.add_row().cells
        write_cell(cells[0], case["case_id"], center=True)
        write_cell(cells[1], detail["scenario"])
        write_cell(cells[2], detail["steps"])
        write_cell(cells[3], detail["expected"])
        write_cell(cells[4], detail["criterion"])
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width(cell, [0.58, 1.03, 1.88, 1.87, 1.24][index])

    # Replace the previous result table/caption and append a richer result section.
    while len(doc.tables) > 1:
        doc.tables[-1]._element.getparent().remove(doc.tables[-1]._element)
    for paragraph in list(doc.paragraphs):
        if paragraph.text.startswith("表 3.1.5.4-1") or paragraph.text.startswith("结果分析与截图插入说明") or paragraph.text.startswith("【此处插入图") or paragraph.text.startswith("图 3.1.5.4-"):
            remove_paragraph(paragraph)

    doc.paragraphs[12].text = "本次实际运行 10 条口语化家政指令，8 条通过、2 条未通过，动作解析准确率为 80.0%，达到整体合格标准。以下表格记录的是未经动作约束纠正的原始模型输出，因此能够反映真实的解析能力和错误类型。"

    add_text(doc, "表 3.1.5.4-1 十个测试用例实际运行结果", bold=True, size=10.5)
    result_table = doc.add_table(rows=1, cols=6)
    result_table.style = table.style
    result_headers = ["用例 ID", "匹配任务", "匹配物体", "匹配操作", "输出序列（原始模型）", "判定结果"]
    for index, text in enumerate(result_headers):
        write_cell(result_table.rows[0].cells[index], text, size=8.5, bold=True, center=True)
        shade(result_table.rows[0].cells[index], "D9EAF7")
    for case in report["cases"]:
        actual = case["actual"]
        resolved = actual["resolved"]
        cells = result_table.add_row().cells
        values = [case["case_id"], resolved["task"], resolved["object"], resolved["operation"], action_lines(actual["predicted_actions"]), "合格" if case["passed"] else "不合格"]
        for index, value in enumerate(values):
            write_cell(cells[index], value, size=8, center=index in {0, 5})
        if not case["passed"]:
            shade(cells[5], "FCE4D6")
    for row in result_table.rows:
        for index, cell in enumerate(row.cells):
            width(cell, [0.56, 0.9, 0.95, 0.85, 2.48, 0.61][index])

    add_text(doc, "结果分析与截图插入说明", bold=True, size=10.5)
    add_text(doc, "从总体结果看，书房桌面收纳、门锁安防检查、遥控器取送、马桶清洁、电池回收、电饭煲烹饪、可视门铃巡检和拐杖取送均完成了任务、物体、操作和原始动作序列的完整匹配。两条未通过用例集中体现了动作方向判别与多物体主从关系识别的改进需求。")
    add_text(doc, "【此处插入图 3.1.5.4-1：十条用例总体测试结果截图】", bold=True, placeholder=True)
    add_text(doc, "图 3.1.5.4-1 应展示测试脚本输出的 8/10 通过统计和 80.0% 准确率，用于说明本次测试的总体结果与合格结论。", size=9.5)
    add_text(doc, "ACT-002 的自然语言指令明确要求关闭阅读灯，系统虽然正确匹配到阅读灯和 turn_off 操作，但原始输出序列出现了 turn_on(reading_lamp)，说明当前模型在设备开关方向的序列生成上仍可能发生反向错误。")
    add_text(doc, "【此处插入图 3.1.5.4-2：ACT-002 阅读灯关闭测试结果截图】", bold=True, placeholder=True)
    add_text(doc, "图 3.1.5.4-2 建议截取输入指令、匹配结果和原始输出序列，重点标注预期 turn_off 与实际 turn_on 的差异。", size=9.5)
    add_text(doc, "ACT-009 中，指令的待操作对象是玩具，收纳箱仅作为放置容器；实际解析却将 storage_box 作为匹配物体并生成围绕收纳箱的动作序列。该结果表明多物体语句需要进一步强化“操作对象—目标容器”的角色区分。")
    add_text(doc, "【此处插入图 3.1.5.4-3：ACT-009 玩具收纳测试结果截图】", bold=True, placeholder=True)
    add_text(doc, "图 3.1.5.4-3 建议截取输入指令、实际匹配物体 storage_box 和原始动作序列，以直观说明玩具与收纳箱主从角色混淆的原因。", size=9.5)

    doc.save(DOC_PATH)


if __name__ == "__main__":
    main()
