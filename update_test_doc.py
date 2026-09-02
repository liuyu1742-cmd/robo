"""Insert real 10-case manual-instruction test evidence into Desktop/3.docx."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


DOC_PATH = Path(r"C:\Users\sjtu101\Desktop\3.docx")
REPORT_PATH = Path(r"C:\RobotProject\RobotProject\datasets\manual_instruction_ten_case_actual_report.json")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def format_cell(cell, *, size: float = 8.5, bold: bool = False, center: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            run.font.size = Pt(size)
            run.bold = bold


def set_cell_text(cell, text: str, *, size: float = 8.5, bold: bool = False, center: bool = False) -> None:
    cell.text = text
    format_cell(cell, size=size, bold=bold, center=center)


def format_paragraph(paragraph, text: str, *, bullet: bool = False) -> None:
    paragraph.text = text
    paragraph.style = "List Paragraph" if bullet else "Normal"
    paragraph.paragraph_format.space_after = Pt(4)
    for run in paragraph.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10.5)


def numbered_actions(actions: list[str]) -> str:
    return "\n".join(f"{index}. {action}" for index, action in enumerate(actions, 1))


def remove_table(table) -> None:
    table._element.getparent().remove(table._element)


def main() -> None:
    report = json.loads(REPORT_PATH.read_text(encoding="utf-8"))
    doc = Document(DOC_PATH)

    # Preserve the original document structure but replace the generic test prose.
    paragraphs = doc.paragraphs
    format_paragraph(paragraphs[1], "在项目虚拟环境中运行 C:\\RobotProject\\RobotProject\\manual_instruction_test.py，加载语义动作生成检查点 action_transformer_project_generation_semantic.pt。", bullet=True)
    format_paragraph(paragraphs[2], "从整理任务、家电综合管理、智慧安防、物品取送、卫浴清洁、垃圾处理和智能烹饪中抽取工况，并使用与基准语句不同的口语化表达作为输入。", bullet=True)
    format_paragraph(paragraphs[3], "逐条记录匹配任务、匹配物体、匹配操作和原始模型输出序列；不以动作约束后的纠正序列替代原始模型输出。", bullet=True)
    format_paragraph(paragraphs[4], "将实际匹配结果和原始输出序列与测试用例中预设的任务、物体、操作及标准动作序列逐项比对。", bullet=True)
    format_paragraph(paragraphs[5], "动作解析准确率 = 原始输出完全匹配的用例数量 ÷ 测试总用例数量 × 100%。", bullet=True)
    format_paragraph(paragraphs[7], "测试环境：Windows 项目虚拟环境；测试入口：manual_instruction_test.py；模型检查点：action_transformer_project_generation_semantic.pt；测试总数：10。")
    format_paragraph(paragraphs[9], "动作解析准确率 = 原始输出中任务、物体、操作及动作序列均与预期一致的用例数量 ÷ 总测试用例数量 × 100%。")
    format_paragraph(paragraphs[10], "合格标准：整体平均准确率≥70%；单条用例须同时满足任务、物体、操作匹配正确，且原始输出动作序列与预期序列完全一致。")
    format_paragraph(paragraphs[12], "本次实际运行 10 条口语化家政指令，8 条通过、2 条未通过，动作解析准确率为 80.0%，达到整体合格标准。未通过原因包括：关闭阅读灯时输出了 turn_on 动作；玩具收纳时将收纳箱误识别为操作对象。")

    case_table = doc.tables[0]
    while len(case_table.rows) > 1:
        case_table._tbl.remove(case_table.rows[-1]._tr)
    headers = ["用例 ID", "测试场景工况", "测试操作步骤", "预期输出结果", "合格判定标准"]
    for index, header in enumerate(headers):
        set_cell_text(case_table.rows[0].cells[index], header, size=9, bold=True, center=True)
        set_cell_shading(case_table.rows[0].cells[index], "D9EAF7")

    for case in report["cases"]:
        expected = case["expected"]
        row = case_table.add_row().cells
        set_cell_text(row[0], case["case_id"], center=True)
        set_cell_text(row[1], f"{case['scenario']}\n输入：{case['input_instruction']}")
        set_cell_text(row[2], "1. 在命令行输入口语化指令；\n2. 启动动作解析入口；\n3. 记录原始模型输出；\n4. 与预期任务、物体、操作和序列比对。")
        set_cell_text(row[3], f"任务：{expected['task']}\n物体：{expected['object']}\n操作：{expected['operation']}\n序列：\n{numbered_actions(expected['actions'])}")
        set_cell_text(row[4], "任务、物体、操作均正确，且原始输出序列逐项完全一致，判定为合格。")
    for row in case_table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, [0.55, 1.35, 1.85, 1.9, 1.0][index])
            format_cell(cell, size=8.3, bold=(row == case_table.rows[0]), center=(index == 0))

    # Remove an old run's result table if this updater is run again.
    while len(doc.tables) > 1:
        remove_table(doc.tables[-1])

    doc.add_paragraph("表 3.1.5.4-1 十个测试用例实际运行结果")
    result_table = doc.add_table(rows=1, cols=6)
    result_table.style = case_table.style
    result_headers = ["用例 ID", "匹配任务", "匹配物体", "匹配操作", "输出序列（原始模型）", "判定结果"]
    for index, header in enumerate(result_headers):
        set_cell_text(result_table.rows[0].cells[index], header, size=8.5, bold=True, center=True)
        set_cell_shading(result_table.rows[0].cells[index], "D9EAF7")
    for case in report["cases"]:
        actual = case["actual"]
        resolved = actual["resolved"]
        row = result_table.add_row().cells
        values = [
            case["case_id"],
            resolved["task"],
            resolved["object"],
            resolved["operation"],
            numbered_actions(actual["predicted_actions"]),
            "合格" if case["passed"] else "不合格",
        ]
        for index, value in enumerate(values):
            set_cell_text(row[index], value, size=8, center=index in {0, 5})
        if not case["passed"]:
            set_cell_shading(row[5], "FCE4D6")
    for row in result_table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, [0.55, 0.9, 0.95, 0.85, 2.5, 0.6][index])
            format_cell(cell, size=8, bold=(row == result_table.rows[0]), center=index in {0, 5})

    doc.save(DOC_PATH)


if __name__ == "__main__":
    main()
